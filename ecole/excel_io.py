"""قالب Excel لإدخال الوضع الحالي للأكاديمية، واستيراده إلى البرنامج.

القالب يحتوي على أوراق: تعليمات، المستويات، المواد، الأساتذة، الأفواج، التلاميذ.
الاستيراد يتحقق من كل البيانات أولاً، ولا يكتب شيئاً إذا وُجد خطأ.
"""

from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .db import DEFAULT_LEVELS, DEFAULT_SUBJECTS

SHEET_HELP = "تعليمات"
SHEET_LEVELS = "المستويات"
SHEET_SUBJECTS = "المواد"
SHEET_TEACHERS = "الأساتذة"
SHEET_GROUPS = "الأفواج"
SHEET_STUDENTS = "التلاميذ"

GROUP_SLOTS = 4  # عدد أعمدة الأفواج لكل تلميذ
MAX_ROWS = {SHEET_LEVELS: 60, SHEET_SUBJECTS: 60, SHEET_TEACHERS: 200,
            SHEET_GROUPS: 300, SHEET_STUDENTS: 2000}

TEACHER_COLS = ["الاسم واللقب *", "الهاتف", "ملاحظات"]
GROUP_COLS = ["رمز الفوج (تلقائي)", "المادة *", "المستوى *", "الأستاذ *", "اسم الفوج *",
              "التوقيت", "القاعة", "الحد الأقصى للتلاميذ"]
STUDENT_COLS = (["الاسم واللقب *", "هاتف التلميذ", "هاتف الولي", "المستوى"]
                + [f"الفوج {i}" for i in range(1, GROUP_SLOTS + 1)]
                + ["الرصيد الحالي (حصص)", "ملاحظات"])

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
AUTO_FILL = PatternFill("solid", start_color="E8EEF7")
BORDER = Border(*(Side(style="thin", color="B8C4D6"),) * 4)


def group_key(subject, level, teacher, name) -> str:
    """نفس صيغة عمود «رمز الفوج» في القالب."""
    return f"{subject} - {level} - {teacher} - {name}"


# ---------------------------------------------------------------- القالب
def _header(ws, headers, widths):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 32


def _body_style(ws, cols, rows, text_cols=()):
    for r in range(2, rows + 2):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=11)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="right")
            if c in text_cols:
                cell.number_format = "@"


def _list_validation(ws, formula, cell_range, title):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "اختر قيمة من القائمة."
    dv.errorTitle = title
    dv.prompt = "اختر من القائمة"
    dv.promptTitle = title
    ws.add_data_validation(dv)
    dv.add(cell_range)


def write_template(path) -> None:
    wb = Workbook()
    help_ws = wb.active
    help_ws.title = SHEET_HELP
    _write_help(help_ws)

    levels = wb.create_sheet(SHEET_LEVELS)
    _header(levels, ["المستوى"], [24])
    _body_style(levels, 1, MAX_ROWS[SHEET_LEVELS])
    for i, name in enumerate(DEFAULT_LEVELS, start=2):
        levels.cell(row=i, column=1, value=name)

    subjects = wb.create_sheet(SHEET_SUBJECTS)
    _header(subjects, ["المادة"], [24])
    _body_style(subjects, 1, MAX_ROWS[SHEET_SUBJECTS])
    for i, name in enumerate(DEFAULT_SUBJECTS, start=2):
        subjects.cell(row=i, column=1, value=name)

    teachers = wb.create_sheet(SHEET_TEACHERS)
    _header(teachers, TEACHER_COLS, [30, 18, 34])
    _body_style(teachers, len(TEACHER_COLS), MAX_ROWS[SHEET_TEACHERS], text_cols=(2,))

    groups = wb.create_sheet(SHEET_GROUPS)
    _header(groups, GROUP_COLS, [48, 18, 14, 26, 12, 24, 10, 12])
    n = MAX_ROWS[SHEET_GROUPS]
    _body_style(groups, len(GROUP_COLS), n)
    for r in range(2, n + 2):
        c = groups.cell(row=r, column=1,
                        value=f'=IF(B{r}="","",B{r}&" - "&C{r}&" - "&D{r}&" - "&E{r})')
        c.fill = AUTO_FILL
        c.font = Font(name=FONT, size=10, color="5B6B80")
    groups["A1"].comment = Comment(
        "يُحسب تلقائياً من المادة والمستوى والأستاذ واسم الفوج. لا تكتب فيه.", "البرنامج")
    _list_validation(groups, f"='{SHEET_SUBJECTS}'!$A$2:$A${MAX_ROWS[SHEET_SUBJECTS] + 1}",
                     f"B2:B{n + 1}", "المادة")
    _list_validation(groups, f"='{SHEET_LEVELS}'!$A$2:$A${MAX_ROWS[SHEET_LEVELS] + 1}",
                     f"C2:C{n + 1}", "المستوى")
    _list_validation(groups, f"='{SHEET_TEACHERS}'!$A$2:$A${MAX_ROWS[SHEET_TEACHERS] + 1}",
                     f"D2:D{n + 1}", "الأستاذ")
    cap = DataValidation(type="whole", operator="between", formula1="0", formula2="500",
                         allow_blank=True)
    cap.error = "اكتب عدداً صحيحاً."
    groups.add_data_validation(cap)
    cap.add(f"H2:H{n + 1}")

    students = wb.create_sheet(SHEET_STUDENTS)
    widths = [30, 16, 16, 14] + [44] * GROUP_SLOTS + [14, 30]
    _header(students, STUDENT_COLS, widths)
    m = MAX_ROWS[SHEET_STUDENTS]
    _body_style(students, len(STUDENT_COLS), m, text_cols=(2, 3))
    _list_validation(students, f"='{SHEET_LEVELS}'!$A$2:$A${MAX_ROWS[SHEET_LEVELS] + 1}",
                     f"D2:D{m + 1}", "المستوى")
    first, last = students.cell(row=1, column=5).column_letter, \
        students.cell(row=1, column=4 + GROUP_SLOTS).column_letter
    _list_validation(students, f"='{SHEET_GROUPS}'!$A$2:$A${n + 1}",
                     f"{first}2:{last}{m + 1}", "الفوج")
    bal_col = students.cell(row=1, column=5 + GROUP_SLOTS).column_letter
    bal = DataValidation(type="whole", operator="between", formula1="-100", formula2="500",
                         allow_blank=True)
    bal.error = "اكتب عدداً صحيحاً (سالب إذا كان على التلميذ دَين)."
    bal.prompt = "موجب: حصص مدفوعة مسبقاً لم تُستهلك. سالب: حصص حضرها ولم يدفعها."
    bal.promptTitle = "الرصيد الحالي"
    students.add_data_validation(bal)
    bal.add(f"{bal_col}2:{bal_col}{m + 1}")
    students[f"{bal_col}1"].comment = Comment(
        "عدد الحصص المتبقية للتلميذ اليوم.\nمثال: دفع باقة 4 وحضر 1 ← اكتب 3.\n"
        "إذا حضر ولم يدفع حصتين ← اكتب -2.\nإذا لا شيء ← اتركه فارغاً.", "البرنامج")

    wb.active = wb.sheetnames.index(SHEET_HELP)
    # يحسب Excel عمود «رمز الفوج» عند فتح الملف.
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def _write_help(ws):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.column_dimensions["A"].width = 3
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 22
    texts = [
        ("title", "قالب إدخال بيانات الأكاديمية"),
        ("text", "هذا الملف لإدخال الوضع الحالي للأكاديمية مرة واحدة، ثم استيراده إلى "
                 "البرنامج من: الإعدادات ← استيراد من Excel."),
        ("head", "ترتيب التعبئة"),
        ("text", "الخطوة 1: «المستويات» و«المواد»: مملوءتان مسبقاً، أضف أو احذف حسب الأكاديمية."),
        ("text", "الخطوة 2: «الأساتذة»: اسم كل أستاذ مرة واحدة فقط."),
        ("text", "الخطوة 3: «الأفواج»: اختر المادة والمستوى والأستاذ من القوائم واكتب اسم الفوج. "
                 "العمود الأول «رمز الفوج» يُملأ تلقائياً، لا تكتب فيه."),
        ("text", "الخطوة 4: «التلاميذ»: اختر أفواج كل تلميذ من القوائم (حتى 4 أفواج) واكتب رصيده "
                 "الحالي من الحصص."),
        ("head", "ملاحظات مهمة"),
        ("text", "الأعمدة المعلّمة بـ * إجبارية، والباقي يمكن تركه فارغاً."),
        ("text", "الرصيد الحالي = عدد الحصص المدفوعة التي لم يحضرها التلميذ بعد. "
                 "اكتبه سالباً إذا حضر حصصاً ولم يدفعها (دَين)، واتركه فارغاً إذا كان صفراً."),
        ("text", "لا تغيّر أسماء الأوراق ولا عناوين الأعمدة."),
        ("text", "يمكن استيراد الملف مرة ثانية بعد إضافة أسطر جديدة: التلميذ الموجود مسبقاً "
                 "(نفس الاسم والمستوى) لا يُكرَّر."),
        ("head", "أمثلة على التعبئة (للتوضيح فقط، لا تُستورد)"),
    ]
    examples = [
        ("ورقة «الأساتذة»", ["الاسم واللقب", "الهاتف", "ملاحظات"],
         ["بن علي محمد", "0661 11 22 33", ""]),
        ("ورقة «الأفواج»", ["المادة", "المستوى", "الأستاذ", "اسم الفوج", "التوقيت", "القاعة"],
         ["رياضيات", "4 متوسط", "بن علي محمد", "فوج 1", "السبت 14:00", "قاعة 1"]),
        ("ورقة «التلاميذ»", ["الاسم واللقب", "هاتف الولي", "المستوى", "الفوج 1",
                            "الرصيد الحالي"],
         ["أحمد بوزيد", "0550 12 34 56", "4 متوسط",
          "رياضيات - 4 متوسط - بن علي محمد - فوج 1", 3]),
    ]
    r = 1
    for kind, value in texts:
        if kind == "head":
            r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c = ws.cell(row=r, column=2, value=value)
        c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        if kind == "title":
            c.font = Font(name=FONT, bold=True, size=16, color="1F3A5F")
            ws.row_dimensions[r].height = 30
        elif kind == "head":
            c.font = Font(name=FONT, bold=True, size=12, color="1F3A5F")
            ws.row_dimensions[r].height = 20
        else:
            c.font = Font(name=FONT, size=11)
            ws.row_dimensions[r].height = 18 if len(value) <= 95 else 33
        r += 1
    for title, headers, sample in examples:
        r += 1
        c = ws.cell(row=r, column=2, value=title)
        c.font = Font(name=FONT, bold=True, size=11)
        r += 1
        for i, (h, v) in enumerate(zip(headers, sample)):
            hc = ws.cell(row=r, column=2 + i, value=h)
            hc.font = Font(name=FONT, bold=True, size=10)
            hc.fill = AUTO_FILL
            hc.border = BORDER
            hc.alignment = Alignment(horizontal="center")
            vc = ws.cell(row=r + 1, column=2 + i, value=v if v != "" else None)
            vc.font = Font(name=FONT, size=10, italic=True, color="5B6B80")
            vc.border = BORDER
            vc.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        ws.row_dimensions[r + 1].height = 30
        r += 2


# --------------------------------------------------------------- الاستيراد
@dataclass
class ImportReport:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    added: dict[str, int] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors

    def count(self, key: str) -> None:
        self.added[key] = self.added.get(key, 0) + 1

    def summary(self) -> str:
        labels = {"levels": "مستويات", "subjects": "مواد", "teachers": "أساتذة",
                  "groups": "أفواج", "students": "تلاميذ", "memberships": "تسجيلات في الأفواج",
                  "balances": "أرصدة افتتاحية"}
        lines = [f"{labels[k]}: {v}" for k, v in self.added.items() if v]
        return "\n".join(lines) or "لم تُضف أي بيانات جديدة."


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _phone(value) -> str:
    text = _text(value)
    # رقم مكتوب كعدد في Excel يفقد الصفر الأول (مثلاً 661112233).
    if text.isdigit() and len(text) == 9:
        text = "0" + text
    return text


def _rows(wb, sheet, report, ncols):
    if sheet not in wb.sheetnames:
        report.errors.append(f"الورقة «{sheet}» غير موجودة في الملف.")
        return []
    ws = wb[sheet]
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=2, max_col=ncols, values_only=True), start=2):
        values = list(row) + [None] * (ncols - len(row))
        if any(_text(v) for v in values):
            out.append((r, values))
    return out


def _names(wb, sheet, report):
    names = []
    for r, values in _rows(wb, sheet, report, 1):
        name = _text(values[0])
        if name in names:
            report.warnings.append(f"«{sheet}» السطر {r}: «{name}» مكرر، تم تجاهله.")
        else:
            names.append(name)
    return names


def import_workbook(store, path, dry_run: bool = False) -> ImportReport:
    """يقرأ القالب، يتحقق منه، ثم يضيف البيانات الجديدة إلى البرنامج.

    إذا وُجد أي خطأ لا يُضاف شيء. dry_run=True للتحقق فقط.
    """
    report = ImportReport()
    try:
        wb = load_workbook(path, read_only=True)
    except Exception as exc:  # ملف تالف أو ليس Excel
        report.errors.append(f"تعذّر فتح الملف: {exc}")
        return report

    level_names = _names(wb, SHEET_LEVELS, report)
    subject_names = _names(wb, SHEET_SUBJECTS, report)
    existing_levels = {l["name"]: l["id"] for l in store.list_levels()}
    existing_subjects = {s["name"]: s["id"] for s in store.list_subjects()}
    known_levels = set(level_names) | set(existing_levels)
    known_subjects = set(subject_names) | set(existing_subjects)

    # الأساتذة
    teachers: dict[str, dict] = {}
    for r, v in _rows(wb, SHEET_TEACHERS, report, len(TEACHER_COLS)):
        name = _text(v[0])
        if not name:
            report.errors.append(f"«{SHEET_TEACHERS}» السطر {r}: اسم الأستاذ فارغ.")
        elif name in teachers:
            report.errors.append(f"«{SHEET_TEACHERS}» السطر {r}: الأستاذ «{name}» مكرر.")
        else:
            teachers[name] = {"full_name": name, "phone": _phone(v[1]), "notes": _text(v[2])}
    existing_teachers = {t["full_name"]: t["id"] for t in store.list_teachers()}
    known_teachers = set(teachers) | set(existing_teachers)

    # الأفواج
    groups: dict[str, dict] = {}
    for r, v in _rows(wb, SHEET_GROUPS, report, len(GROUP_COLS)):
        subject, level, teacher, name = (_text(x) for x in v[1:5])
        if not any((subject, level, teacher, name)):
            continue  # سطر لا يحتوي إلا على صيغة رمز الفوج
        where = f"«{SHEET_GROUPS}» السطر {r}"
        missing = [label for label, val in (("المادة", subject), ("المستوى", level),
                                            ("الأستاذ", teacher), ("اسم الفوج", name)) if not val]
        if missing:
            report.errors.append(f"{where}: ناقص: {'، '.join(missing)}.")
            continue
        if subject not in known_subjects:
            report.errors.append(f"{where}: المادة «{subject}» غير موجودة في ورقة المواد.")
        if level not in known_levels:
            report.errors.append(f"{where}: المستوى «{level}» غير موجود في ورقة المستويات.")
        if teacher not in known_teachers:
            report.errors.append(f"{where}: الأستاذ «{teacher}» غير موجود في ورقة الأساتذة.")
        key = group_key(subject, level, teacher, name)
        if key in groups:
            report.errors.append(f"{where}: الفوج «{key}» مكرر.")
            continue
        capacity = v[7]
        try:
            capacity = int(capacity) if _text(capacity) else 0
        except (TypeError, ValueError):
            report.errors.append(f"{where}: الحد الأقصى يجب أن يكون عدداً.")
            capacity = 0
        groups[key] = {"subject": subject, "level": level, "teacher": teacher, "name": name,
                       "schedule": _text(v[5]), "room": _text(v[6]), "capacity": capacity}
    existing_groups = {group_key(g["subject_name"], g["level_name"], g["teacher_name"],
                                 g["name"]): g["id"] for g in store.list_groups()}
    known_groups = set(groups) | set(existing_groups)

    # التلاميذ
    students = []
    seen_students = set()
    existing_students = {(s["full_name"], s["level_name"] or "") for s in store.list_students()}
    for r, v in _rows(wb, SHEET_STUDENTS, report, len(STUDENT_COLS)):
        where = f"«{SHEET_STUDENTS}» السطر {r}"
        name, level = _text(v[0]), _text(v[3])
        if not name:
            report.errors.append(f"{where}: اسم التلميذ فارغ.")
            continue
        if level and level not in known_levels:
            report.errors.append(f"{where}: المستوى «{level}» غير موجود في ورقة المستويات.")
        group_keys = []
        for x in v[4:4 + GROUP_SLOTS]:
            key = _text(x)
            if not key:
                continue
            if key not in known_groups:
                report.errors.append(f"{where}: الفوج «{key}» غير موجود في ورقة الأفواج.")
            elif key not in group_keys:
                group_keys.append(key)
        raw_balance = v[4 + GROUP_SLOTS]
        try:
            balance = int(float(_text(raw_balance))) if _text(raw_balance) else 0
        except ValueError:
            report.errors.append(f"{where}: الرصيد «{_text(raw_balance)}» يجب أن يكون عدداً.")
            balance = 0
        ident = (name, level)
        if ident in seen_students:
            report.errors.append(f"{where}: التلميذ «{name}» ({level}) مكرر في الملف.")
            continue
        seen_students.add(ident)
        if ident in existing_students:
            report.warnings.append(f"{where}: «{name}» موجود مسبقاً في البرنامج، لم يُستورد.")
            continue
        students.append({"full_name": name, "phone": _phone(v[1]),
                         "parent_phone": _phone(v[2]), "level": level,
                         "notes": _text(v[5 + GROUP_SLOTS]), "groups": group_keys,
                         "balance": balance})
    wb.close()

    if report.errors or dry_run:
        return report

    # -------------------------------------------------------- الكتابة
    for name in level_names:
        if name not in existing_levels:
            existing_levels[name] = store.save_level(None, name)
            report.count("levels")
    for name in subject_names:
        if name not in existing_subjects:
            existing_subjects[name] = store.save_subject(None, name)
            report.count("subjects")
    for name, data in teachers.items():
        if name in existing_teachers:
            report.warnings.append(f"الأستاذ «{name}» موجود مسبقاً، لم يُضف مرة ثانية.")
            continue
        existing_teachers[name] = store.save_teacher(data)
        report.count("teachers")
    for key, g in groups.items():
        if key in existing_groups:
            continue
        existing_groups[key] = store.save_group({
            "name": g["name"], "teacher_id": existing_teachers[g["teacher"]],
            "subject_id": existing_subjects[g["subject"]],
            "level_id": existing_levels[g["level"]], "schedule": g["schedule"],
            "room": g["room"], "capacity": g["capacity"],
        })
        report.count("groups")
    for s in students:
        sid = store.save_student({
            "full_name": s["full_name"], "phone": s["phone"], "parent_phone": s["parent_phone"],
            "level_id": existing_levels.get(s["level"]), "notes": s["notes"],
        })
        report.count("students")
        for key in s["groups"]:
            store.add_group_member(existing_groups[key], sid)
            report.count("memberships")
        if s["balance"]:
            store.add_balance_adjustment(sid, s["balance"], "رصيد افتتاحي (استيراد Excel)")
            report.count("balances")
    return report
