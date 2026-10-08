import pytest
from openpyxl import load_workbook

from ecole import db, excel_io
from ecole.excel_io import SHEET_GROUPS, SHEET_STUDENTS, SHEET_TEACHERS
from ecole.services import Store


@pytest.fixture
def store():
    s = Store(db.connect(":memory:"))
    s.authenticate("admin", "admin")
    return s


def filled_template(tmp_path, students=None):
    path = tmp_path / "template.xlsx"
    excel_io.write_template(path)
    wb = load_workbook(path)
    t = wb[SHEET_TEACHERS]
    t.append(["بن علي محمد", "0661112233"])
    t["A3"], t["B3"] = "سعيدي فاطمة", 770445566  # هاتف مكتوب كعدد
    g = wb[SHEET_GROUPS]
    for r, row in ((2, ["رياضيات", "4 متوسط", "بن علي محمد", "فوج 1", "السبت 14:00", "ق1", 20]),
                   (3, ["فيزياء", "4 متوسط", "سعيدي فاطمة", "فوج 1", "", "", None])):
        for c, value in enumerate(row, start=2):
            g.cell(row=r, column=c, value=value)
    math = "رياضيات - 4 متوسط - بن علي محمد - فوج 1"
    phys = "فيزياء - 4 متوسط - سعيدي فاطمة - فوج 1"
    s = wb[SHEET_STUDENTS]
    rows = students or [
        ["أحمد بوزيد", "", "0550123456", "4 متوسط", math, phys, None, None, 3],
        ["سارة حمدي", "", "", "4 متوسط", math, None, None, None, -2],
        ["يوسف شريف", "", "", "4 متوسط", None, None, None, None, None],
    ]
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            s.cell(row=r, column=c, value=value)
    wb.save(path)
    return path


def test_template_has_all_sheets(tmp_path):
    path = tmp_path / "t.xlsx"
    excel_io.write_template(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["تعليمات", "المستويات", "المواد", "الأساتذة", "الأفواج", "التلاميذ"]
    assert wb[SHEET_GROUPS]["A2"].value.startswith("=IF(")


def test_import_full(store, tmp_path):
    report = excel_io.import_workbook(store, filled_template(tmp_path))
    assert report.ok, report.errors
    assert report.added["teachers"] == 2
    assert report.added["groups"] == 2
    assert report.added["students"] == 3
    assert report.added["memberships"] == 3
    students = {s["full_name"]: s for s in store.list_students()}
    assert students["أحمد بوزيد"]["balance"] == 3
    assert students["سارة حمدي"]["balance"] == -2
    assert students["يوسف شريف"]["balance"] == 0
    assert students["أحمد بوزيد"]["parent_phone"] == "0550123456"
    phones = {t["full_name"]: t["phone"] for t in store.list_teachers()}
    assert phones["سعيدي فاطمة"] == "0770445566"
    # الرصيد الافتتاحي ليس مبلغاً مقبوضاً
    assert store.academy_summary("2000-01-01", "2999-12-31")["cash"] == 0


def test_reimport_skips_existing(store, tmp_path):
    path = filled_template(tmp_path)
    excel_io.import_workbook(store, path)
    report = excel_io.import_workbook(store, path)
    assert report.ok
    assert report.added.get("students", 0) == 0
    assert len(store.list_students()) == 3
    assert store.list_students(search="أحمد")[0]["balance"] == 3


def test_errors_block_import(store, tmp_path):
    path = filled_template(tmp_path, students=[
        ["", "", "", "4 متوسط"],
        ["تلميذ", "", "", "مستوى خيالي", "فوج غير موجود", None, None, None, "ثلاثة"],
    ])
    report = excel_io.import_workbook(store, path)
    assert not report.ok
    assert len(report.errors) == 4
    assert store.list_students() == []
    assert store.list_teachers() == []


def test_opening_balance_used_by_attendance(store, tmp_path):
    excel_io.import_workbook(store, filled_template(tmp_path))
    ahmed = store.list_students(search="أحمد")[0]["id"]
    group = store.student_groups(ahmed)[0]["id"]
    store.save_attendance(store.get_or_create_lesson(group, "2026-10-08"), {ahmed: True})
    assert store.student_balance(ahmed) == 2
    labels = [e["label"] for e in store.student_statement(ahmed)]
    assert labels[0].startswith("تعديل الرصيد")
